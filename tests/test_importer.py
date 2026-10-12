"""Only synthetic workbooks: real client data must never enter versioned tests."""

from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
import json
import sqlite3
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest

from importer import _a_number, import_workbook
from storage import connect_db


def workbook(tmp_path, sheets, filename="synthetic.xlsx"):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    path = tmp_path / filename
    book.save(path)
    return path


def query(database, sql):
    connection = connect_db(database)
    try:
        return [dict(row) for row in connection.execute(sql)]
    finally:
        connection.close()


def test_matches_all_name_tokens_and_identifiers_without_conflicting_merge(tmp_path):
    path = workbook(tmp_path, {
        "RD Status": [
            ["Attorney", "Client Last Name", "Client First Name", "A#", "JCATS"],
            ["Another attorney", "Pérez López", "Ana María", "123-456-789", "O26-100"],
            ["Mario", "Pérez López", "Ana María", "999-456-789", "O26-100"],
            ["Mario", "García", "Lucía", "A123456788", "O26101"],
        ],
        "Client Calls": [
            ["Date", "Client Name", "Reason for Calling"],
            [datetime(2026, 6, 23), "GARCIA, LUCIA", "Información"],
            [datetime(2026, 6, 24), "Ana María Pérez López", "Nombre ambiguo"],
            [datetime(2026, 6, 25), "Lucía", "Nombre parcial"],
        ],
        "Documents Received": [
            ["Date Received", "Client Name", "A #", "Document(s) Type"],
            [datetime(2026, 6, 26), "Nombre alternativo completo", "123456788", "Carta"],
        ],
    })
    database = tmp_path / "cases.sqlite3"
    report = import_workbook(path, database)
    assert report["sources_imported"] == 7
    assert report["cases_created"] == 5  # Three anchored, one ambiguous, one partial.
    cases = query(database, "SELECT * FROM cases")
    assert len({case["id"] for case in cases if case["a_number"] in {"123456789", "999456789"}}) == 2
    source_ids = query(database, "SELECT case_id,worksheet,row_number FROM sources")
    lucia = next(case["id"] for case in cases if case["a_number"] == "123456788")
    assert sum(row["case_id"] == lucia for row in source_ids) == 3
    assert report["identity_issues"]


def test_sections_duplicate_columns_merged_cells_and_precise_skips(tmp_path):
    path = workbook(tmp_path, {"Clients Info": [
        ["Former Clients"],
        ["Client's Name", "A-NUMBER", "JCATS", "email", "email", None],
        ["Ana Pérez", "123456789", "O26101", "one@example.test", "two@example.test", "Original extra"],
        [None, None, None, "Missing identity"],
        ["Active Clients"],
        ["#A Number", "Client Name", None, "JCAT", "email"],
        ["123456788", "José", "García", "O26102", "three@example.test"],
    ], "Affirmative": [
        ["Attorney", "Active/Closed?", "Client Last Name", "Client First Name", "Pending Assignments"],
        ["Mario", "Activo", "Pérez", "Ana", "Solicitar documento"],
        ["Mario", None, "López", None, None],
    ]})
    book = load_workbook(path)
    for column in (2, 4, 5):
        book["Affirmative"].merge_cells(start_row=2, start_column=column, end_row=3, end_column=column)
    book.save(path)
    database = tmp_path / "cases.sqlite3"
    report = import_workbook(path, database)
    assert report["sources_imported"] == 4
    assert report["tasks_created"] == 2
    sources = query(database, "SELECT worksheet,row_number,data FROM sources")
    first = json.loads(next(row["data"] for row in sources if row["worksheet"] == "Clients Info" and row["row_number"] == 3))
    assert first["email [D]"] == "one@example.test"
    assert first["email [E]"] == "two@example.test"
    assert first["Columna F"] == "Original extra"
    assert "José García" in {case["name"] for case in query(database, "SELECT name FROM cases")}
    merged = json.loads(next(row["data"] for row in sources if row["worksheet"] == "Affirmative" and row["row_number"] == 3))
    assert merged["Client First Name"] == "Ana"
    assert merged["Pending Assignments"] == "Solicitar documento"
    skips = {(item["worksheet"], item["row_number"]): item["reason"] for item in report["skipped_rows"]}
    assert skips[("Clients Info", 2)] == "encabezado"
    assert "sin nombre" in skips[("Clients Info", 4)]
    assert "sección" in skips[("Clients Info", 5)]


def test_dates_explicit_tasks_completed_assignments_and_active_case(tmp_path):
    path = workbook(tmp_path, {
        "RD Status": [
            ["Attorney", "Client Last Name", "Client First Name", "A#", "MCH Date", "Next Due Date", "Notes"],
            ["[Closed] Mario", "Pérez", "Ana", "123456789", "Granted Asylum", "urgent tomorrow", datetime(2026, 3, 8)],
        ],
        "Affirmative": [
            ["Attorney", "Active/Closed?", "Client Last Name", "Client First Name", "A#", "Pending Assignments"],
            ["Mario", "Activo", "Pérez", "Ana", "123456789", "Solicitar firmas"],
        ],
        "Assigments": [
            ["First Name", "Last Name", "A Number", "Assignment", "Assigned Clerk", "Deadline", "Date Completed"],
            ["Ana", "Pérez", "123456789", "Preparar escrito", "Asistente", datetime(2026, 7, 4), datetime(2026, 7, 3)],
            ["Ana", "Pérez", "123456789", "Revisar información", "Asistente", "7/8/2026", "Done"],
            ["Ana", "Pérez", "123456789", "Llamar", "Asistente", None, None],
        ],
        "Client Calls": [
            ["Date", "Client Name", "Reason for Calling"],
            [datetime(2026, 6, 23), "Ana Pérez", "Llamada"],
        ],
        "Documents Received": [
            ["Date Received", "Client Name", "Document(s) Type"],
            ["6/7/2026", "Ana Pérez", "Carta"],
            ["6/23/2026", "Ana Pérez", "Otra carta"],
        ],
    })
    database = tmp_path / "cases.sqlite3"
    report = import_workbook(path, database, author="Lucía")
    assert report["tasks_created"] == 4
    tasks = {row["title"]: row for row in query(database, "SELECT * FROM tasks")}
    assert tasks["Preparar escrito"]["due_date"] == "2026-07-04"
    assert tasks["Preparar escrito"]["status"] == "completada"
    assert tasks["Revisar información"]["status"] == "completada"
    assert tasks["Revisar información"]["due_date"] == ""
    assert "7/8/2026" in tasks["Revisar información"]["description"]
    assert tasks["Llamar"]["due_date"] == ""
    assert all(task["author"] == "Lucía" for task in tasks.values())
    case = query(database, "SELECT * FROM cases")[0]
    assert case["status"] == "activo"
    assert case["hearing_date"] == ""
    source = json.loads(query(database, "SELECT data FROM sources WHERE worksheet='RD Status'")[0]["data"])
    assert source["Notes"] == "2026-03-08"
    assert source["Next Due Date"] == "urgent tomorrow"
    activities = query(database, "SELECT * FROM activities ORDER BY id")
    assert [row["occurred_on"] for row in activities] == ["2026-06-23", "", "2026-06-23"]
    assert all(row["author"] == "Lucía" for row in activities)


def test_same_and_revised_workbook_preserve_duplicate_rows_and_manual_records(tmp_path):
    path = workbook(tmp_path, {"Client Calls": [
        ["Date", "Client Name", "Reason for Calling"],
        [datetime(2026, 6, 23), "Ana Pérez", "Información"],
        [datetime(2026, 6, 23), "Ana Pérez", "Información"],
    ]})
    database = tmp_path / "cases.sqlite3"
    report = import_workbook(path, database)
    assert report["sources_imported"] == report["activities_created"] == 2
    repeated = import_workbook(path, database)
    assert repeated["already_imported"]
    assert repeated["cases_created"] == repeated["sources_imported"] == repeated["activities_created"] == 0
    connection = connect_db(database)
    case_id = connection.execute("SELECT id FROM cases").fetchone()[0]
    with connection:
        connection.execute("UPDATE cases SET notes='Nota manual',country='Manual',status='cerrado' WHERE id=?", (case_id,))
        connection.execute("INSERT INTO activities (case_id,type,title,author,created_at) VALUES (?,?,?,?,?)",
                           (case_id, "nota", "Manual", "Mario", "2026-10-12T00:00:00+00:00"))
    connection.close()
    book = load_workbook(path)
    book.active.insert_rows(2)
    book.active.append([datetime(2026, 6, 24), "Ana Pérez", "Seguimiento"])
    revised = tmp_path / "revised.xlsx"
    book.save(revised)
    report = import_workbook(revised, database)
    assert not report["already_imported"]
    assert report["cases_created"] == 0
    assert report["sources_imported"] == report["activities_created"] == 1
    assert len(query(database, "SELECT * FROM sources")) == 3
    assert len(query(database, "SELECT * FROM activities")) == 4
    case = query(database, "SELECT * FROM cases")[0]
    assert (case["notes"], case["country"], case["status"]) == ("Nota manual", "Manual", "cerrado")


def test_missing_names_with_valid_identifier_are_preserved(tmp_path):
    path = workbook(tmp_path, {"Cases": [
        ["Client Name", "A#", "JCATS", "Other"],
        [None, "123456789", None, "Keep all fields"],
        ["Partial", None, None, "Keep partial name"],
        [None, "invalid", None, "Skip invalid identifier precisely"],
    ]})
    database = tmp_path / "cases.sqlite3"
    report = import_workbook(path, database)
    assert report["sources_imported"] == 2
    assert report["cases_created"] == 2
    assert any(item["row_number"] == 4 and "sin nombre" in item["reason"] for item in report["skipped_rows"])


def test_reject_invalid_and_excessively_large_archives(tmp_path, monkeypatch):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="XLSX"):
        import_workbook(path, tmp_path / "cases.sqlite3")
    with ZipFile(path, "w") as archive:
        archive.writestr("large.xml", b"x" * 101)
    monkeypatch.setattr("importer.MAX_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(ValueError, match="descomprimidos"):
        import_workbook(path, tmp_path / "cases.sqlite3")


@pytest.mark.parametrize("value,expected", [
    ("A-241-138-481", "241138481"), ("A- 245-224-965", "245224965"),
    ("012-345-678", "012345678"), (12345678, "012345678"),
    ("N/A", ""), ("Phone 123456789", ""), ("1234", ""),
])
def test_a_number_normalization(value, expected):
    assert _a_number(value) == expected


def test_concurrent_reimport_is_idempotent(tmp_path):
    path = workbook(tmp_path, {"Client Calls": [
        ["Date", "Client Name", "Reason for Calling"],
        [datetime(2026, 6, 23), "Ana Pérez", "Información"],
    ]})
    database = tmp_path / "cases.sqlite3"
    # Initialize before the independent requests, as Flask does at startup.
    from storage import init_db
    init_db(database)
    with ThreadPoolExecutor(max_workers=2) as pool:
        reports = list(pool.map(lambda _: import_workbook(path, database), range(2)))
    assert sum(report["cases_created"] for report in reports) == 1
    assert sum(report["sources_imported"] for report in reports) == 1
    assert sorted(report["already_imported"] for report in reports) == [False, True]
    assert len(query(database, "SELECT * FROM activities")) == 1
