"""Conservative, repeatable import of case workbooks without executing cell content."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from storage import connect_db, init_db


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_SHEETS = 50
MAX_ROWS = 50_000
MAX_CELLS = 1_000_000


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return " ".join(re.findall(r"[a-z0-9]+", "".join(
        c for c in text.casefold() if not unicodedata.combining(c))))


def _present(value):
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _text(value):
    return "" if value is None else str(value).strip()


def _original_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _date(value):
    """Only actual dates and individually unambiguous whole-cell strings."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:T\d{2}:\d{2}:\d{2}(?:[.+Z\d:-]*)?)?", text):
        try:
            return date.fromisoformat(text[:10]).isoformat()
        except ValueError:
            return ""
    # E.g. 6/23/2026 has one possible ordering. 6/7/2026 remains original text.
    match = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", text)
    if match:
        first, second, year = map(int, match.groups())
        if first > 12 >= second:
            day, month = first, second
        elif second > 12 >= first:
            month, day = first, second
        else:
            return ""
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            return ""
    return ""


def _a_number(value):
    # USCIS permits seven to nine digits; normalize to a nine-digit A-number.
    if isinstance(value, int) and not isinstance(value, bool):
        value = str(value)
    elif isinstance(value, float) and value.is_integer():
        value = str(int(value))
    text = _text(value).upper()
    if re.fullmatch(r"(?:A[\s-]*)?[0-9][0-9\s-]*", text):
        digits = re.sub(r"\D", "", text)
        if 7 <= len(digits) <= 9:
            return digits.zfill(9)
    return ""


def _reference(value):
    text = _text(value).upper()
    if _normalize(text) in {"", "n a", "na", "none", "no", "unknown", "pending"}:
        return ""
    return re.sub(r"[^A-Z0-9]", "", text)


def _name_key(name):
    tokens = _normalize(name).split()
    if len(tokens) < 2 or _normalize(name) in {"no name", "not available", "client name"}:
        return ""
    return " ".join(sorted(tokens))


ALIASES = {
    "last": {"client last name", "last name", "apellido", "apellidos"},
    "first": {"client first name", "first name", "nombre"},
    "name": {"client name", "client s name", "name", "nombre completo"},
    "a": {"a", "a number", "a numero", "numero a"},
    "reference": {"jcats", "jcat"},
    "owner": {"attorney", "attorney assigned"},
    "status": {"active closed", "status"},
    "country": {"respondent s country", "country", "pais"},
    "court": {"court", "district filed", "ij", "judge assigned"},
    "hearing": {"next hearing date", "mch date", "hearing date", "next hearing"},
    "assignment": {"assignment", "pending assignments"},
    "assignee": {"assigned clerk"},
    "due": {"deadline"},
    "completed": {"date completed"},
    "assigned_on": {"date assigned"},
}
LOOKUP = {alias: key for key, aliases in ALIASES.items() for alias in aliases}


def _field(label):
    # Duplicate source headers include their Excel letter while keeping the original label.
    return LOOKUP.get(_normalize(re.sub(r" \[[A-Z]+\]$", "", str(label))), "")


def _fields(data):
    result = {}
    for label, value in data.items():
        key = _field(label)
        if key and _present(value) and key not in result:
            result[key] = value
    return result


def _client_name(fields, data, worksheet):
    name = _text(fields.get("name")) or " ".join(
        part for part in (_text(fields.get("first")), _text(fields.get("last"))) if part)
    # Later Clients Info sections split the displayed name between B and C,
    # although their merged header only labels B. Keep C as its original field too.
    if _normalize(worksheet) == "clients info" and fields.get("name") and _present(data.get("Columna C")):
        name = " ".join(filter(None, (name, _text(data["Columna C"]))))
    if _normalize(name) in {"n a", "na", "none", "unknown", "no name", "pending", "tbd", "not available"}:
        return ""
    return name


def _is_header(values):
    fields = {_field(value) for value in values if _present(value)} - {""}
    return bool(fields & {"name", "first", "last"}) and len(fields) >= 1 and (
        len(fields) >= 2 or any(_normalize(value) in ALIASES["name"] for value in values)
    )


def _labels(values):
    raw = [_text(value) if _present(value) else f"Columna {get_column_letter(i)}"
           for i, value in enumerate(values, 1)]
    counts = Counter(raw)
    return [f"{label} [{get_column_letter(i)}]" if counts[label] > 1 else label
            for i, label in enumerate(raw, 1)]


@dataclass
class SourceRow:
    worksheet: str
    row_number: int
    data: dict
    fields: dict
    name: str
    a_number: str
    reference: str
    fingerprint: str
    case_id: int | None = None


def _open_workbook(path):
    path = Path(path)
    if not path.is_file():
        raise ValueError("No se encontró el archivo Excel.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("El archivo supera el límite de 20 MB.")
    try:
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("El Excel supera el límite de 100 MB descomprimidos.")
        book = load_workbook(path, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("El archivo no es un libro Excel XLSX válido.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("No se pudo leer el libro Excel XLSX.") from exc
    if len(book.worksheets) > MAX_SHEETS:
        book.close()
        raise ValueError("El Excel supera el límite de 50 hojas.")
    if any(sheet.max_row > MAX_ROWS for sheet in book):
        book.close()
        raise ValueError("Una hoja supera el límite de 50.000 filas.")
    if sum(sheet.max_row * sheet.max_column for sheet in book) > MAX_CELLS:
        book.close()
        raise ValueError("El Excel supera el límite de 1.000.000 de celdas.")
    return book


def _read_rows(book, report):
    result = []
    duplicate_counts = Counter()
    for sheet in book:
        info = {"name": sheet.title, "rows_imported": 0, "rows_skipped": 0}
        report["sheets"].append(info)
        labels = None
        section = ""
        section_rows = {merged.min_row for merged in sheet.merged_cells.ranges
                        if merged.min_row == merged.max_row and merged.min_col == 1
                        and merged.max_col >= sheet.max_column - 1}
        inherited = {}
        for cell_range in sheet.merged_cells.ranges:
            value = sheet.cell(cell_range.min_row, cell_range.min_col).value
            for row in range(cell_range.min_row, cell_range.max_row + 1):
                for column in range(cell_range.min_col, cell_range.max_col + 1):
                    inherited[(row, column)] = value
        for cells in sheet.iter_rows():
            row_number = cells[0].row
            original = [cell.value for cell in cells]
            if not any(_present(value) for value in original):
                continue
            values = [inherited.get((row_number, i), value) for i, value in enumerate(original, 1)]
            if _is_header(original):
                labels = _labels(original)
                reason = "encabezado"
            elif labels is None:
                section = " | ".join(_text(value) for value in original if _present(value))
                reason = "sección o título anterior al encabezado"
            elif (_normalize(sheet.title) == "clients info" and row_number in section_rows) or (
                    sum(_present(value) for value in original) == 1 and (
                    _normalize(next(value for value in original if _present(value))) in {
                        "former clients", "active clients", "current clients", "new clients", "clients info"})):
                section = " | ".join(_text(value) for value in original if _present(value))
                reason = "título de sección"
            else:
                data = {label: _original_value(value) for label, value in zip(labels, values) if _present(value)}
                if section:
                    data["Sección original"] = section
                fields = _fields(data)
                name = _client_name(fields, data, sheet.title)
                a_number = _a_number(fields.get("a"))
                reference = _reference(fields.get("reference"))
                if name or a_number or reference:
                    canonical = json.dumps([sheet.title, data], ensure_ascii=False, sort_keys=True, separators=(",", ":"))
                    duplicate_counts[canonical] += 1
                    fingerprint = hashlib.sha256((canonical + f"#{duplicate_counts[canonical]}").encode()).hexdigest()
                    result.append(SourceRow(sheet.title, row_number, data, fields, name, a_number, reference, fingerprint))
                    continue
                reason = "sin nombre de cliente ni A-number/JCATS válidos"
            info["rows_skipped"] += 1
            report["rows_skipped"] += 1
            report["skipped_rows"].append({"worksheet": sheet.title, "row_number": row_number, "reason": reason})
    return result


class Identities:
    def __init__(self, connection, report, timestamp):
        self.connection, self.report, self.timestamp = connection, report, timestamp
        self.cases = {row["id"]: dict(row) for row in connection.execute("SELECT * FROM cases")}
        self.by_a, self.by_ref, self.by_name = defaultdict(set), defaultdict(set), defaultdict(set)
        self.case_refs = defaultdict(set)
        for case_id, case in self.cases.items():
            self.add_alias(case_id, case["name"], _a_number(case["a_number"]), _reference(case["reference"]))
        for row in connection.execute("SELECT case_id,worksheet,data FROM sources"):
            try:
                data = json.loads(row["data"])
                fields = _fields(data)
            except (ValueError, TypeError):
                continue
            name = _client_name(fields, data, row["worksheet"])
            self.add_alias(row["case_id"], name, _a_number(fields.get("a")), _reference(fields.get("reference")))

    def add_alias(self, case_id, name, a_number, reference):
        if a_number:
            self.by_a[a_number].add(case_id)
        if reference:
            self.by_ref[reference].add(case_id)
            self.case_refs[case_id].add(reference)
        if _name_key(name):
            self.by_name[_name_key(name)].add(case_id)

    def compatible(self, case_id, row, compare_reference=True):
        existing_a = _a_number(self.cases[case_id]["a_number"])
        if existing_a and row.a_number and existing_a != row.a_number:
            return False
        refs = self.case_refs[case_id]
        return not (compare_reference and refs and row.reference and row.reference not in refs)

    def resolve(self, row):
        candidates = self.by_a[row.a_number] if row.a_number else set()
        if len(candidates) == 1:
            case_id = next(iter(candidates))
        else:
            refs = self.by_ref[row.reference] if row.reference else set()
            compatible_refs = {i for i in refs if self.compatible(i, row, compare_reference=False)}
            if len(compatible_refs) == 1:
                case_id = next(iter(compatible_refs))
            else:
                names = self.by_name[_name_key(row.name)] if _name_key(row.name) else set()
                if len(names) == 1 and self.compatible(next(iter(names)), row):
                    case_id = next(iter(names))
                else:
                    if candidates or refs or names:
                        self.report["identity_issues"].append({
                            "worksheet": row.worksheet, "row_number": row.row_number,
                            "reason": "identificadores incompatibles o nombre ambiguo; se conservó un caso separado",
                        })
                    case_id = self.create(row)
        self.add_alias(case_id, row.name, row.a_number, row.reference)
        self.fill_empty(case_id, row)
        return case_id

    def create(self, row):
        name = row.name or (f"Cliente A-{row.a_number}" if row.a_number else f"Cliente {row.reference}")
        owner = re.sub(r"\[[^]]*\]", "", _text(row.fields.get("owner"))).strip() or "Mario"
        notes = "\n".join(f"{label}: {value}" for label, value in row.data.items()
                          if _normalize(label) in {"notes", "case progress notes", "case progress notes appointments", "current status"})
        cursor = self.connection.execute(
            "INSERT INTO cases (name,a_number,reference,category,status,owner,country,court,hearing_date,notes,created_at,updated_at) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
            (name, row.a_number, row.reference, row.worksheet, "activo", owner, "", "", "", notes, self.timestamp, self.timestamp),
        )
        case_id = cursor.lastrowid
        self.cases[case_id] = dict(self.connection.execute("SELECT * FROM cases WHERE id=?", (case_id,)).fetchone())
        self.report["cases_created"] += 1
        return case_id

    def fill_empty(self, case_id, row):
        case = self.cases[case_id]
        owner = re.sub(r"\[[^]]*\]", "", _text(row.fields.get("owner"))).strip() or "Mario"
        changes = {"a_number": row.a_number, "reference": row.reference,
                   "country": _text(row.fields.get("country")), "court": _text(row.fields.get("court")),
                   "hearing_date": _date(row.fields.get("hearing")), "owner": owner}
        # Imports enrich empty fields but never overwrite later manual edits.
        changes = {key: value for key, value in changes.items() if value and not case.get(key)}
        if changes:
            self.connection.execute("UPDATE cases SET " + ", ".join(f"{key}=?" for key in changes) +
                                    ", updated_at=? WHERE id=?", (*changes.values(), self.timestamp, case_id))
            case.update(changes)


def _description(row):
    return "\n".join(f"{label}: {value}" for label, value in row.data.items())


def _data_value(row, *labels):
    normalized = {_normalize(label): value for label, value in row.data.items()}
    for label in labels:
        value = normalized.get(_normalize(label))
        if _present(value):
            return value
    return ""


def _seed_activity(connection, row, report, timestamp, author):
    sheet = _normalize(row.worksheet)
    if sheet == "client calls":
        kind, title = "llamada", "Llamada del cliente"
        occurred = _date(_data_value(row, "Date", "Date/Fecha", "l", "Date Received", "Fecha"))
    elif sheet == "documents received":
        kind = "documento"
        document = _text(_data_value(row, "Document(s) Type", "Document Type"))
        title = "Documentos recibidos" + (f": {document}" if document else "")
        occurred = _date(_data_value(row, "Date Received", "Date", "Fecha"))
    else:
        return
    cursor = connection.execute(
        "INSERT OR IGNORE INTO activities (case_id,type,title,description,occurred_on,author,created_at,source_key) "
        "VALUES (?,?,?,?,?,?,?,?)",
        (row.case_id, kind, title, _description(row), occurred, author, timestamp, row.fingerprint + ":activity"),
    )
    report["activities_created"] += cursor.rowcount


def _seed_task(connection, row, report, timestamp, author):
    assignment = _text(row.fields.get("assignment"))
    if not assignment or _normalize(assignment) in {"none", "n a", "na", "no", "ninguna", "ninguno"}:
        return
    # Only explicit assignment columns create tasks, never free-form legal notes.
    if _normalize(row.worksheet) not in {"assigments", "assignments", "affirmative"}:
        return
    completion = row.fields.get("completed")
    done = _present(completion) and _normalize(completion) not in {"n a", "na", "no", "pending", "pendiente", "not completed"}
    completed_on = _date(completion)
    completed_at = completed_on + "T00:00:00+00:00" if done and completed_on else ""
    cursor = connection.execute(
        "INSERT OR IGNORE INTO tasks (case_id,title,description,assignee,due_date,priority,status,author,updated_by,"
        "completed_at,created_at,updated_at,source_key) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (row.case_id, assignment, _description(row), _text(row.fields.get("assignee")) or "Mario",
         _date(row.fields.get("due")), "media", "completada" if done else "pendiente", author,
         author, completed_at, timestamp, timestamp, row.fingerprint + ":task"),
    )
    report["tasks_created"] += cursor.rowcount


def _status_evidence(worksheet, data):
    fields = _fields(data)
    text = ""
    if _normalize(worksheet) == "rd status":
        text = _text(fields.get("owner"))
    elif _normalize(worksheet) in {"case status", "affirmative"}:
        text = _text(fields.get("status"))
    if not text:
        return ""
    tokens = set(_normalize(text).split())
    return "cerrado" if tokens & {"closed", "cerrado", "cerrada", "inactivo"} else "activo"


def _finalize_statuses(connection, new_ids):
    # A closed proceeding is not enough to close a client with another active source.
    evidence = defaultdict(set)
    for source in connection.execute("SELECT case_id,worksheet,data FROM sources"):
        if source["case_id"] in new_ids:
            status = _status_evidence(source["worksheet"], json.loads(source["data"]))
            if status:
                evidence[source["case_id"]].add(status)
    for case_id, statuses in evidence.items():
        if statuses == {"cerrado"}:
            connection.execute("UPDATE cases SET status='cerrado' WHERE id=?", (case_id,))


def _report_identity_quality(rows, report):
    references = defaultdict(list)
    for row in rows:
        if row.reference and row.a_number:
            references[row.reference].append(row)
        if len(_normalize(row.name).split()) < 2:
            report["incomplete_names"].append({
                "worksheet": row.worksheet, "row_number": row.row_number,
                "reason": "nombre ausente o parcial; se conservó la fuente sin fusionar por nombre",
            })
    for sources in references.values():
        if len({row.a_number for row in sources}) > 1:
            first = sources[0]
            report["identity_issues"].append({
                "worksheet": first.worksheet, "row_number": first.row_number,
                "reason": "JCATS compartido por A-number distintos; se conservaron los casos separados",
                "related_rows": [{"worksheet": row.worksheet, "row_number": row.row_number} for row in sources],
            })


def import_workbook(path, db_path, author="Importación Excel"):
    """Import a local XLSX atomically; unchanged and moved rows are idempotent."""
    author = _text(author) or "Importación Excel"
    book = _open_workbook(path)
    file_hash = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    init_db(db_path)
    connection = connect_db(db_path)
    try:
        # Serialize imports before reading their identities/hash; simultaneous
        # uploads of the same workbook must not create two copies of each case.
        connection.execute("BEGIN IMMEDIATE")
        previous = connection.execute("SELECT report FROM import_runs WHERE file_hash=?", (file_hash,)).fetchone()
        if previous:
            report = json.loads(previous["report"])
            return {**report, "cases_created": 0, "sources_imported": 0, "tasks_created": 0,
                    "activities_created": 0, "already_imported": True,
                    "source_count": connection.execute("SELECT COUNT(*) FROM sources").fetchone()[0]}
        report = {"cases_created": 0, "sources_imported": 0, "tasks_created": 0,
                  "activities_created": 0, "rows_skipped": 0, "sheets": [],
                  "already_imported": False, "skipped_rows": [], "identity_issues": [],
                  "incomplete_names": [], "author": author}
        rows = _read_rows(book, report)
        if not rows:
            raise ValueError("No se encontraron filas de clientes con encabezados reconocibles en el Excel.")
        _report_identity_quality(rows, report)
        timestamp = _now()
        with connection:
            identities = Identities(connection, report, timestamp)
            old_ids = set(identities.cases)
            # Anchored identities go first, preventing an ambiguous name-only row from
            # merging two distinct people whose identifiers appear later in the book.
            rows.sort(key=lambda row: (not bool(row.a_number), not bool(row.reference)))
            existing = {row["fingerprint"]: row["case_id"] for row in connection.execute("SELECT fingerprint,case_id FROM sources")}
            sheet_info = {info["name"]: info for info in report["sheets"]}
            for row in rows:
                if row.fingerprint in existing:
                    row.case_id = existing[row.fingerprint]
                    continue
                row.case_id = identities.resolve(row)
                connection.execute(
                    "INSERT INTO sources (case_id,worksheet,row_number,data,fingerprint,imported_at) VALUES (?,?,?,?,?,?)",
                    (row.case_id, row.worksheet, row.row_number, json.dumps(row.data, ensure_ascii=False), row.fingerprint, timestamp),
                )
                existing[row.fingerprint] = row.case_id
                report["sources_imported"] += 1
                sheet_info[row.worksheet]["rows_imported"] += 1
                _seed_activity(connection, row, report, timestamp, author)
                _seed_task(connection, row, report, timestamp, author)
            _finalize_statuses(connection, set(identities.cases) - old_ids)
            report["source_count"] = connection.execute("SELECT COUNT(*) FROM sources").fetchone()[0]
            connection.execute("INSERT INTO import_runs (filename,file_hash,imported_at,report) VALUES (?,?,?,?)",
                               (Path(path).name, file_hash, timestamp, json.dumps(report, ensure_ascii=False)))
        return report
    finally:
        book.close()
        connection.close()


def main():
    parser = argparse.ArgumentParser(description="Importar casos desde un libro XLSX.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--database", type=Path, default=Path("instance/cases.sqlite3"))
    args = parser.parse_args()
    try:
        print(json.dumps(import_workbook(args.workbook, args.database), ensure_ascii=False, indent=2))
    except ValueError as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    main()
